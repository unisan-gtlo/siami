import io
from xml.sax.saxutils import escape

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import FileResponse, HttpResponse
from django.shortcuts import get_object_or_404, render
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from apps.ami_assessment.models import Pengisian
from apps.ami_core.models import Siklus
from apps.ami_de.models import DeDaftarTilik, DePenugasan
from apps.ami_temuan.models import Temuan

DOCX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


def _docx_response(document, filename):
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type=DOCX_CONTENT_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _docx_signature_table(document, kiri_label, kiri_nama, kanan_label=None, kanan_nama=None):
    table = document.add_table(rows=4, cols=2 if kanan_label else 1)
    table.rows[0].cells[0].paragraphs[0].add_run(kiri_label).bold = True
    table.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.rows[3].cells[0].paragraphs[0].add_run(f'( {kiri_nama} )')
    table.rows[3].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if kanan_label:
        table.rows[0].cells[1].paragraphs[0].add_run(kanan_label).bold = True
        table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        table.rows[3].cells[1].paragraphs[0].add_run(f'( {kanan_nama} )')
        table.rows[3].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def _get_user_ami(request):
    return getattr(request.user, 'ami_profile', None)


def _pengisian_milik_user(user_ami, siklus):
    """Cari Pengisian (bukan bikin baru -- laporan hanya untuk yang sudah
    ada isinya) sesuai cakupan user: prodi > fakultas-UPM > universitas-LP3M."""
    if user_ami is None or siklus is None:
        return None
    if user_ami.prodi_id:
        return Pengisian.objects.filter(siklus=siklus, cakupan='prodi', prodi=user_ami.prodi).first()
    if user_ami.is_upm and user_ami.fakultas_id:
        return Pengisian.objects.filter(siklus=siklus, cakupan='fakultas', fakultas=user_ami.fakultas).first()
    if user_ami.is_lp3m or user_ami.is_pimpinan:
        return Pengisian.objects.filter(siklus=siklus, cakupan='universitas').first()
    return None


@login_required
def laporan_list(request):
    return render(request, 'ami_pelaporan/laporan_list.html', {'active_tab': 'laporan'})


@login_required
def laporan_temuan_pdf(request):
    user_ami = _get_user_ami(request)
    siklus = Siklus.objects.filter(is_current=True).first()
    pengisian = _pengisian_milik_user(user_ami, siklus)
    if pengisian is None:
        messages.error(request, 'Akun Anda belum terhubung ke profil AMI (prodi/fakultas), atau belum ada data pengisian.')
        return render(request, 'ami_pelaporan/no_profile.html', {'active_tab': 'laporan'})

    temuan_qs = Temuan.objects.filter(siklus=siklus, pengisian=pengisian).select_related('butir__master_standar').order_by('klasifikasi')
    praktik_baik_qs = temuan_qs.filter(klasifikasi__in=['SESUAI', 'BP'], layak_replikasi=True)
    temuan_ketidaksesuaian_qs = temuan_qs.exclude(klasifikasi__in=['SESUAI', 'BP'])

    tim_auditor = list(
        DePenugasan.objects.filter(pengisian=pengisian).select_related('auditor__user').order_by(
            '-role_dalam_tim',
        ),
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph('TEMUAN AUDIT', styles['Title']),
    ]
    if siklus and siklus.regulasi_acuan:
        elements.append(Paragraph(f'Acuan: {siklus.regulasi_acuan.nama_pendek}', styles['Normal']))
    elements.append(Spacer(1, 0.3*cm))

    cell_style = styles['Normal'].clone('cell')
    cell_style.fontSize = 8

    def cell(text):
        return Paragraph(escape(str(text)), cell_style)

    info_rows = [
        [cell('Auditee'), cell(f': {pengisian.subjek}'), cell('Siklus'), cell(f': {siklus}')],
        [cell('Fak/Prodi'), cell(f': {pengisian.prodi.fakultas if pengisian.prodi_id else pengisian.get_cakupan_display()}'),
         cell('Tanggal'), cell(f': {timezone.now().date()}')],
    ]
    for a in tim_auditor:
        info_rows.append([cell(a.get_role_dalam_tim_display()), cell(f': {a.auditor}'), cell(''), cell('')])
    info_table = Table(info_rows, colWidths=[3*cm, 6*cm, 3*cm, 6*cm])
    info_table.setStyle(TableStyle([('FONTSIZE', (0, 0), (-1, -1), 8)]))
    elements.append(info_table)
    elements.append(Spacer(1, 0.5*cm))

    elements.append(Paragraph('Temuan Audit (Ketidaksesuaian)', styles['Heading3']))
    data = [['No', 'Klasifikasi', 'Standar', 'Temuan Audit', 'Tenggat', 'Status']]
    for i, t in enumerate(temuan_ketidaksesuaian_qs, start=1):
        data.append([
            str(i),
            t.get_klasifikasi_display(),
            cell(t.butir.master_standar.kode if t.butir_id and t.butir.master_standar_id else '-'),
            cell(f'{t.judul} — {t.deskripsi_problem[:150]}'),
            str(t.tenggat_tindak_lanjut or '-'),
            t.get_status_display(),
        ])
    if len(data) == 1:
        elements.append(Paragraph('Tidak ada temuan ketidaksesuaian.', styles['Normal']))
    else:
        table = Table(data, colWidths=[1*cm, 2.3*cm, 2*cm, 8*cm, 2*cm, 2.2*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0E5A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E8F2F8')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FAFCFE')]),
        ]))
        elements.append(table)
    elements.append(Spacer(1, 0.5*cm))

    elements.append(Paragraph('Praktik Baik (Temuan Positif)', styles['Heading3']))
    data2 = [['No', 'Aspek/Bidang', 'Kelebihan']]
    for i, t in enumerate(praktik_baik_qs, start=1):
        data2.append([
            str(i),
            cell(t.butir.master_standar.kode if t.butir_id and t.butir.master_standar_id else '-'),
            cell(f'{t.judul} — {t.deskripsi_problem[:200]}'),
        ])
    if len(data2) == 1:
        elements.append(Paragraph('Belum ada praktik baik teridentifikasi.', styles['Normal']))
    else:
        table2 = Table(data2, colWidths=[1*cm, 3*cm, 13.5*cm])
        table2.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#27AE60')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E8F2F8')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5FFFC')]),
        ]))
        elements.append(table2)
    elements.append(Spacer(1, 1.2*cm))

    ttd_auditee = pengisian.operator or 'Auditee'
    ttd_auditor = tim_auditor[0].auditor if tim_auditor else 'Auditor'
    ttd_table = Table(
        [
            ['Menyetujui,', ''],
            ['Auditee', 'Auditor'],
            ['', ''],
            ['', ''],
            [f'( {ttd_auditee} )', f'( {ttd_auditor} )'],
        ],
        colWidths=[8.5*cm, 8.5*cm],
    )
    ttd_table.setStyle(TableStyle([
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('LINEBELOW', (0, 3), (0, 3), 0.5, colors.black),
        ('LINEBELOW', (1, 3), (1, 3), 0.5, colors.black),
    ]))
    elements.append(ttd_table)

    doc.build(elements)
    buffer.seek(0)

    kode_subjek = pengisian.prodi.kode if pengisian.prodi_id else pengisian.cakupan
    filename = f'laporan-temuan-{kode_subjek}-siklus{siklus.no_siklus if siklus else "x"}.pdf'
    return FileResponse(buffer, as_attachment=True, filename=filename, content_type='application/pdf')


@login_required
def laporan_daftar_tilik_pdf(request, penugasan_id):
    user_ami = _get_user_ami(request)
    penugasan = get_object_or_404(DePenugasan, pk=penugasan_id)

    is_monitor = request.user.is_superuser or (user_ami and (user_ami.is_lp3m or user_ami.is_pimpinan))
    if not is_monitor and (user_ami is None or penugasan.auditor_id != user_ami.id):
        messages.error(request, 'Anda tidak punya akses ke Daftar Tilik penugasan ini.')
        return render(request, 'ami_pelaporan/no_profile.html', {'active_tab': 'laporan'})

    tilik_qs = DeDaftarTilik.objects.filter(penugasan=penugasan).select_related(
        'butir__master_standar', 'butir__standar',
    ).order_by('no_urut')

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    cell_style = styles['Normal'].clone('cell')
    cell_style.fontSize = 8

    def cell(text):
        return Paragraph(escape(str(text)), cell_style)

    elements = [Paragraph('DAFTAR TILIK', styles['Title'])]

    info_rows = [
        [cell('Auditee'), cell(f': {penugasan.pengisian.subjek}'), cell('Tanggal'), cell(f': {penugasan.tgl_mulai_de}')],
        [cell('Auditor'), cell(f': {penugasan.auditor}'), cell('SK No'), cell(f': {penugasan.sk_no or "-"}')],
    ]
    info_table = Table(info_rows, colWidths=[3*cm, 6*cm, 3*cm, 6*cm])
    info_table.setStyle(TableStyle([('FONTSIZE', (0, 0), (-1, -1), 8)]))
    elements.append(info_table)
    elements.append(Spacer(1, 0.5*cm))

    data = [['No', 'Standar', 'Pertanyaan/Cek', 'Hasil', 'Catatan']]
    for dt in tilik_qs:
        if dt.butir and dt.butir.master_standar_id:
            standar_label = dt.butir.master_standar.kode
        elif dt.butir and dt.butir.standar_id:
            standar_label = dt.butir.standar.nama_pendek
        else:
            standar_label = '-'
        data.append([
            str(dt.no_urut or ''),
            cell(standar_label),
            cell(dt.deskripsi_tilik),
            dt.get_status_visitasi_display(),
            cell(dt.catatan_visitasi or ''),
        ])

    if len(data) == 1:
        elements.append(Paragraph('Belum ada Daftar Tilik untuk penugasan ini.', styles['Normal']))
    else:
        table = Table(data, colWidths=[1*cm, 2.5*cm, 8*cm, 2.5*cm, 3*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0E5A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E8F2F8')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FAFCFE')]),
        ]))
        elements.append(table)

    elements.append(Spacer(1, 1.2*cm))
    ttd_table = Table(
        [['Auditor'], [''], [''], [f'( {penugasan.auditor} )']],
        colWidths=[8.5*cm],
    )
    ttd_table.setStyle(TableStyle([
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('LINEBELOW', (0, 2), (0, 2), 0.5, colors.black),
    ]))
    elements.append(ttd_table)

    doc.build(elements)
    buffer.seek(0)

    filename = f'daftar-tilik-penugasan-{penugasan.id}.pdf'
    return FileResponse(buffer, as_attachment=True, filename=filename, content_type='application/pdf')


@login_required
def laporan_naratif_pdf(request):
    """Laporan AMI naratif -- Proses & Metode / Temuan Audit / Praktik Baik /
    Rekomendasi / Simpulan -- mengikuti format "Contoh Form Laporan AMI 1.docx"
    yang disediakan LP3M. Cuma bagian yang punya data riil yang ditulis;
    bagian tanpa data ditandai jujur, bukan diisi teks generik."""
    user_ami = _get_user_ami(request)
    siklus = Siklus.objects.filter(is_current=True).first()
    pengisian = _pengisian_milik_user(user_ami, siklus)
    if pengisian is None:
        messages.error(request, 'Akun Anda belum terhubung ke profil AMI (prodi/fakultas), atau belum ada data pengisian.')
        return render(request, 'ami_pelaporan/no_profile.html', {'active_tab': 'laporan'})

    temuan_qs = Temuan.objects.filter(siklus=siklus, pengisian=pengisian)
    ketidaksesuaian_qs = temuan_qs.exclude(klasifikasi__in=['SESUAI', 'BP'])
    praktik_baik_qs = temuan_qs.filter(klasifikasi__in=['SESUAI', 'BP'], layak_replikasi=True)
    tim_auditor = list(DePenugasan.objects.filter(pengisian=pengisian).select_related('auditor__user'))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    body = styles['Normal'].clone('body')
    body.spaceAfter = 8

    elements = [
        Paragraph('LAPORAN AUDIT MUTU INTERNAL', styles['Title']),
        Paragraph(f'Auditee: {pengisian.subjek}', body),
        Paragraph(f'Auditor Internal: {", ".join(str(a.auditor) for a in tim_auditor) or "-"}', body),
        Paragraph(f'Hari/Tanggal AMI: {timezone.now().date()}', body),
        Paragraph(f'Lingkup Audit: {siklus}', body),
        Spacer(1, 0.4*cm),
    ]

    elements.append(Paragraph('PROSES &amp; METODE AUDIT', styles['Heading3']))
    elements.append(Paragraph(
        f'Audit dilaksanakan melalui Desk Evaluasi berbasis {pengisian.total_butir} butir instrumen '
        f'({pengisian.butir_terisi} terisi auditee) pada Siklus {siklus.no_siklus if siklus else "-"}, '
        f'dinilai oleh {len(tim_auditor)} auditor yang ditugaskan.', body,
    ))

    elements.append(Paragraph('TEMUAN AUDIT', styles['Heading3']))
    if ketidaksesuaian_qs.exists():
        for t in ketidaksesuaian_qs:
            elements.append(Paragraph(f'&bull; [{t.get_klasifikasi_display()}] {t.judul} — {t.deskripsi_problem}', body))
    else:
        elements.append(Paragraph('Tidak ada temuan ketidaksesuaian pada siklus berjalan.', body))

    elements.append(Paragraph('PRAKTIK BAIK', styles['Heading3']))
    if praktik_baik_qs.exists():
        for t in praktik_baik_qs:
            elements.append(Paragraph(f'&bull; {t.judul} — {t.deskripsi_problem}', body))
    else:
        elements.append(Paragraph('Belum ada praktik baik yang diidentifikasi layak direplikasi.', body))

    elements.append(Paragraph('REKOMENDASI', styles['Heading3']))
    rekomendasi = [t for t in ketidaksesuaian_qs if t.klasifikasi in ('KTS_MAYOR', 'KTB')]
    if rekomendasi:
        for t in rekomendasi:
            elements.append(Paragraph(f'&bull; Tindak lanjuti segera "{t.judul}" — tenggat {t.tenggat_tindak_lanjut or "belum ditetapkan"}.', body))
    else:
        elements.append(Paragraph('Tidak ada rekomendasi eskalasi prioritas tinggi pada siklus berjalan.', body))

    elements.append(Paragraph('SIMPULAN', styles['Heading3']))
    total = temuan_qs.count()
    elements.append(Paragraph(
        f'Total {total} temuan tercatat pada siklus ini: {ketidaksesuaian_qs.count()} ketidaksesuaian, '
        f'{praktik_baik_qs.count()} praktik baik. Progres self-assessment {pengisian.persentase_progress}% '
        f'({pengisian.butir_terisi}/{pengisian.total_butir} butir).', body,
    ))
    elements.append(Spacer(1, 1.2*cm))

    ttd_style = styles['Normal'].clone('ttd')
    ttd_style.alignment = 1
    ttd_table = Table(
        [[Paragraph('Tanda tangan<br/>Auditor Internal 1', ttd_style), Paragraph('Tanda tangan<br/>Auditor Internal 2', ttd_style)],
         [''], [''],
         [f'( {tim_auditor[0].auditor} )' if len(tim_auditor) > 0 else '( _____________ )',
          f'( {tim_auditor[1].auditor} )' if len(tim_auditor) > 1 else '( _____________ )']],
        colWidths=[8.5*cm, 8.5*cm],
    )
    ttd_table.setStyle(TableStyle([
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('LINEBELOW', (0, 2), (0, 2), 0.5, colors.black),
        ('LINEBELOW', (1, 2), (1, 2), 0.5, colors.black),
    ]))
    elements.append(ttd_table)

    doc.build(elements)
    buffer.seek(0)

    kode_subjek = pengisian.prodi.kode if pengisian.prodi_id else pengisian.cakupan
    filename = f'laporan-ami-naratif-{kode_subjek}-siklus{siklus.no_siklus if siklus else "x"}.pdf'
    return FileResponse(buffer, as_attachment=True, filename=filename, content_type='application/pdf')


# ===================================================================
# Versi DOCX -- sama isinya dengan versi PDF di atas, tapi format Word
# supaya LP3M/auditor bisa mengedit lanjut (tambah narasi, koreksi,
# dsb) sebelum jadi dokumen final. Query data disengaja diulang persis
# sama seperti versi PDF, bukan dibagikan lewat helper -- dua library
# rendering (reportlab vs python-docx) cukup beda sehingga penyatuan
# logika tampilan justru bikin kode lebih susah dibaca.
# ===================================================================

@login_required
def laporan_temuan_docx(request):
    user_ami = _get_user_ami(request)
    siklus = Siklus.objects.filter(is_current=True).first()
    pengisian = _pengisian_milik_user(user_ami, siklus)
    if pengisian is None:
        messages.error(request, 'Akun Anda belum terhubung ke profil AMI (prodi/fakultas), atau belum ada data pengisian.')
        return render(request, 'ami_pelaporan/no_profile.html', {'active_tab': 'laporan'})

    temuan_qs = Temuan.objects.filter(siklus=siklus, pengisian=pengisian).select_related('butir__master_standar').order_by('klasifikasi')
    praktik_baik_qs = temuan_qs.filter(klasifikasi__in=['SESUAI', 'BP'], layak_replikasi=True)
    ketidaksesuaian_qs = temuan_qs.exclude(klasifikasi__in=['SESUAI', 'BP'])
    tim_auditor = list(DePenugasan.objects.filter(pengisian=pengisian).select_related('auditor__user').order_by('-role_dalam_tim'))

    document = Document()
    title = document.add_heading('TEMUAN AUDIT', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if siklus and siklus.regulasi_acuan:
        document.add_paragraph(f'Acuan: {siklus.regulasi_acuan.nama_pendek}')

    info = document.add_table(rows=2 + len(tim_auditor), cols=4)
    info.rows[0].cells[0].text = 'Auditee'
    info.rows[0].cells[1].text = f': {pengisian.subjek}'
    info.rows[0].cells[2].text = 'Siklus'
    info.rows[0].cells[3].text = f': {siklus}'
    info.rows[1].cells[0].text = 'Fak/Prodi'
    info.rows[1].cells[1].text = f': {pengisian.prodi.fakultas if pengisian.prodi_id else pengisian.get_cakupan_display()}'
    info.rows[1].cells[2].text = 'Tanggal'
    info.rows[1].cells[3].text = f': {timezone.now().date()}'
    for i, a in enumerate(tim_auditor, start=2):
        info.rows[i].cells[0].text = a.get_role_dalam_tim_display()
        info.rows[i].cells[1].text = f': {a.auditor}'
    document.add_paragraph()

    document.add_heading('Temuan Audit (Ketidaksesuaian)', level=2)
    if ketidaksesuaian_qs.exists():
        table = document.add_table(rows=1, cols=6)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        for i, h in enumerate(['No', 'Klasifikasi', 'Standar', 'Temuan Audit', 'Tenggat', 'Status']):
            hdr[i].text = h
            hdr[i].paragraphs[0].runs[0].bold = True
        for i, t in enumerate(ketidaksesuaian_qs, start=1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t.get_klasifikasi_display()
            row[2].text = t.butir.master_standar.kode if t.butir_id and t.butir.master_standar_id else '-'
            row[3].text = f'{t.judul} — {t.deskripsi_problem}'
            row[4].text = str(t.tenggat_tindak_lanjut or '-')
            row[5].text = t.get_status_display()
    else:
        document.add_paragraph('Tidak ada temuan ketidaksesuaian.')
    document.add_paragraph()

    document.add_heading('Praktik Baik (Temuan Positif)', level=2)
    if praktik_baik_qs.exists():
        table2 = document.add_table(rows=1, cols=3)
        table2.style = 'Table Grid'
        hdr2 = table2.rows[0].cells
        for i, h in enumerate(['No', 'Aspek/Bidang', 'Kelebihan']):
            hdr2[i].text = h
            hdr2[i].paragraphs[0].runs[0].bold = True
        for i, t in enumerate(praktik_baik_qs, start=1):
            row = table2.add_row().cells
            row[0].text = str(i)
            row[1].text = t.butir.master_standar.kode if t.butir_id and t.butir.master_standar_id else '-'
            row[2].text = f'{t.judul} — {t.deskripsi_problem}'
    else:
        document.add_paragraph('Belum ada praktik baik teridentifikasi.')
    document.add_paragraph()
    document.add_paragraph()

    ttd_auditee = pengisian.operator or 'Auditee'
    ttd_auditor = tim_auditor[0].auditor if tim_auditor else 'Auditor'
    _docx_signature_table(document, 'Auditee', ttd_auditee, 'Auditor', ttd_auditor)

    kode_subjek = pengisian.prodi.kode if pengisian.prodi_id else pengisian.cakupan
    filename = f'laporan-temuan-{kode_subjek}-siklus{siklus.no_siklus if siklus else "x"}.docx'
    return _docx_response(document, filename)


@login_required
def laporan_naratif_docx(request):
    user_ami = _get_user_ami(request)
    siklus = Siklus.objects.filter(is_current=True).first()
    pengisian = _pengisian_milik_user(user_ami, siklus)
    if pengisian is None:
        messages.error(request, 'Akun Anda belum terhubung ke profil AMI (prodi/fakultas), atau belum ada data pengisian.')
        return render(request, 'ami_pelaporan/no_profile.html', {'active_tab': 'laporan'})

    temuan_qs = Temuan.objects.filter(siklus=siklus, pengisian=pengisian)
    ketidaksesuaian_qs = temuan_qs.exclude(klasifikasi__in=['SESUAI', 'BP'])
    praktik_baik_qs = temuan_qs.filter(klasifikasi__in=['SESUAI', 'BP'], layak_replikasi=True)
    tim_auditor = list(DePenugasan.objects.filter(pengisian=pengisian).select_related('auditor__user'))

    document = Document()
    title = document.add_heading('LAPORAN AUDIT MUTU INTERNAL', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f'Auditee: {pengisian.subjek}')
    document.add_paragraph(f'Auditor Internal: {", ".join(str(a.auditor) for a in tim_auditor) or "-"}')
    document.add_paragraph(f'Hari/Tanggal AMI: {timezone.now().date()}')
    document.add_paragraph(f'Lingkup Audit: {siklus}')
    document.add_paragraph()

    document.add_heading('PROSES & METODE AUDIT', level=2)
    document.add_paragraph(
        f'Audit dilaksanakan melalui Desk Evaluasi berbasis {pengisian.total_butir} butir instrumen '
        f'({pengisian.butir_terisi} terisi auditee) pada Siklus {siklus.no_siklus if siklus else "-"}, '
        f'dinilai oleh {len(tim_auditor)} auditor yang ditugaskan.',
    )

    document.add_heading('TEMUAN AUDIT', level=2)
    if ketidaksesuaian_qs.exists():
        for t in ketidaksesuaian_qs:
            document.add_paragraph(f'[{t.get_klasifikasi_display()}] {t.judul} — {t.deskripsi_problem}', style='List Bullet')
    else:
        document.add_paragraph('Tidak ada temuan ketidaksesuaian pada siklus berjalan.')

    document.add_heading('PRAKTIK BAIK', level=2)
    if praktik_baik_qs.exists():
        for t in praktik_baik_qs:
            document.add_paragraph(f'{t.judul} — {t.deskripsi_problem}', style='List Bullet')
    else:
        document.add_paragraph('Belum ada praktik baik yang diidentifikasi layak direplikasi.')

    document.add_heading('REKOMENDASI', level=2)
    rekomendasi = [t for t in ketidaksesuaian_qs if t.klasifikasi in ('KTS_MAYOR', 'KTB')]
    if rekomendasi:
        for t in rekomendasi:
            document.add_paragraph(
                f'Tindak lanjuti segera "{t.judul}" — tenggat {t.tenggat_tindak_lanjut or "belum ditetapkan"}.',
                style='List Bullet',
            )
    else:
        document.add_paragraph('Tidak ada rekomendasi eskalasi prioritas tinggi pada siklus berjalan.')

    document.add_heading('SIMPULAN', level=2)
    total = temuan_qs.count()
    document.add_paragraph(
        f'Total {total} temuan tercatat pada siklus ini: {ketidaksesuaian_qs.count()} ketidaksesuaian, '
        f'{praktik_baik_qs.count()} praktik baik. Progres self-assessment {pengisian.persentase_progress}% '
        f'({pengisian.butir_terisi}/{pengisian.total_butir} butir).',
    )
    document.add_paragraph()
    document.add_paragraph()

    ttd1 = tim_auditor[0].auditor if len(tim_auditor) > 0 else '_____________'
    ttd2 = tim_auditor[1].auditor if len(tim_auditor) > 1 else '_____________'
    _docx_signature_table(document, 'Auditor Internal 1', ttd1, 'Auditor Internal 2', ttd2)

    kode_subjek = pengisian.prodi.kode if pengisian.prodi_id else pengisian.cakupan
    filename = f'laporan-ami-naratif-{kode_subjek}-siklus{siklus.no_siklus if siklus else "x"}.docx'
    return _docx_response(document, filename)


@login_required
def laporan_daftar_tilik_docx(request, penugasan_id):
    user_ami = _get_user_ami(request)
    penugasan = get_object_or_404(DePenugasan, pk=penugasan_id)

    is_monitor = request.user.is_superuser or (user_ami and (user_ami.is_lp3m or user_ami.is_pimpinan))
    if not is_monitor and (user_ami is None or penugasan.auditor_id != user_ami.id):
        messages.error(request, 'Anda tidak punya akses ke Daftar Tilik penugasan ini.')
        return render(request, 'ami_pelaporan/no_profile.html', {'active_tab': 'laporan'})

    tilik_qs = DeDaftarTilik.objects.filter(penugasan=penugasan).select_related(
        'butir__master_standar', 'butir__standar',
    ).order_by('no_urut')

    document = Document()
    title = document.add_heading('DAFTAR TILIK', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = document.add_table(rows=2, cols=4)
    info.rows[0].cells[0].text = 'Auditee'
    info.rows[0].cells[1].text = f': {penugasan.pengisian.subjek}'
    info.rows[0].cells[2].text = 'Tanggal'
    info.rows[0].cells[3].text = f': {penugasan.tgl_mulai_de}'
    info.rows[1].cells[0].text = 'Auditor'
    info.rows[1].cells[1].text = f': {penugasan.auditor}'
    info.rows[1].cells[2].text = 'SK No'
    info.rows[1].cells[3].text = f': {penugasan.sk_no or "-"}'
    document.add_paragraph()

    if tilik_qs.exists():
        table = document.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        for i, h in enumerate(['No', 'Standar', 'Pertanyaan/Cek', 'Hasil', 'Catatan']):
            hdr[i].text = h
            hdr[i].paragraphs[0].runs[0].bold = True
        for dt in tilik_qs:
            if dt.butir and dt.butir.master_standar_id:
                standar_label = dt.butir.master_standar.kode
            elif dt.butir and dt.butir.standar_id:
                standar_label = dt.butir.standar.nama_pendek
            else:
                standar_label = '-'
            row = table.add_row().cells
            row[0].text = str(dt.no_urut or '')
            row[1].text = standar_label
            row[2].text = dt.deskripsi_tilik
            row[3].text = dt.get_status_visitasi_display()
            row[4].text = dt.catatan_visitasi or ''
    else:
        document.add_paragraph('Belum ada Daftar Tilik untuk penugasan ini.')
    document.add_paragraph()
    document.add_paragraph()

    _docx_signature_table(document, 'Auditor', penugasan.auditor)

    filename = f'daftar-tilik-penugasan-{penugasan.id}.docx'
    return _docx_response(document, filename)
