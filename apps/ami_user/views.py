from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import models
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from apps.ami_master.models import Fakultas

from .models import UserAmi

# Referensi hak akses per role -- deskriptif, mencerminkan pengecekan yang
# sudah ditegakkan di masing-masing view (is_auditor_de, prodi_id, is_lp3m,
# dst). BUKAN ACL yang bisa dikonfigurasi dari sini.
PERMISSION_MATRIX = {
    'auditor_de': {
        'label': 'Auditor DE',
        'akses': [
            ('1. Identitas Prodi', 'View only'),
            ('2. Self-Assessment', 'View only'),
            ('3. Upload Bukti', 'View only'),
            ('4. Desk Evaluasi', 'Full Edit'),
            ('5. Visitasi', 'No access'),
            ('6. Temuan dan PTK', 'View only'),
            ('7. Tindak Lanjut', 'No access'),
            ('8. RTM', 'No access'),
            ('9. Dashboard', 'No access'),
            ('10. Pelaporan', 'No access'),
            ('11. Manajemen User', 'No access'),
            ('12. Pengarsipan', 'No access'),
        ],
    },
    'auditor_visitasi': {
        'label': 'Auditor Visitasi',
        'akses': [
            ('1. Identitas Prodi', 'View only'),
            ('2. Self-Assessment', 'View only'),
            ('3. Upload Bukti', 'View only'),
            ('4. Desk Evaluasi', 'View only'),
            ('5. Visitasi', 'Full Edit'),
            ('6. Temuan dan PTK', 'View only'),
            ('7. Tindak Lanjut', 'View only'),
            ('8. RTM', 'No access'),
            ('9. Dashboard', 'No access'),
            ('10. Pelaporan', 'No access'),
            ('11. Manajemen User', 'No access'),
            ('12. Pengarsipan', 'No access'),
        ],
    },
    'auditee': {
        'label': 'Auditee (Prodi)',
        'akses': [
            ('1. Identitas Prodi', 'Edit sebagian'),
            ('2. Self-Assessment', 'Full Edit'),
            ('3. Upload Bukti', 'Full Edit'),
            ('4. Desk Evaluasi', 'No access'),
            ('5. Visitasi', 'View + Konfirmasi'),
            ('6. Temuan dan PTK', 'View only'),
            ('7. Tindak Lanjut', 'Update progress (jika PIC)'),
            ('8. RTM', 'No access'),
            ('9. Dashboard', 'No access'),
            ('10. Pelaporan', 'Generate own'),
            ('11. Manajemen User', 'No access'),
            ('12. Pengarsipan', 'No access'),
        ],
    },
    'lp3m': {
        'label': 'LP3M / Pimpinan',
        'akses': [
            ('1. Identitas Prodi', 'View all'),
            ('2. Self-Assessment', 'View all'),
            ('3. Upload Bukti', 'View all'),
            ('4. Desk Evaluasi', 'View all'),
            ('5. Visitasi', 'View all'),
            ('6. Temuan dan PTK', 'View all'),
            ('7. Tindak Lanjut', 'View all'),
            ('8. RTM', 'Full Edit'),
            ('9. Dashboard', 'Full Edit'),
            ('10. Pelaporan', 'Full Edit'),
            ('11. Manajemen User', 'Full Edit'),
            ('12. Pengarsipan', 'View all'),
        ],
    },
}


def _get_user_ami(request):
    return getattr(request.user, 'ami_profile', None)


def _is_pengawas(request):
    user_ami = _get_user_ami(request)
    return request.user.is_superuser or (user_ami and (user_ami.is_lp3m or user_ami.is_pimpinan))


def ada_konflik_kepentingan(user_ami):
    """True kalau auditor ini punya penugasan DE untuk prodi di fakultas
    sendiri -- crossover mandatory (prinsip independensi Pasal 67 ayat 2
    Permendiktisaintek 39/2025, sudah dicatat di catatan khusus role
    auditor). Dihitung dari penugasan riil, bukan diasumsikan."""
    from apps.ami_de.models import DePenugasan

    if not user_ami.fakultas_id:
        return False
    return DePenugasan.objects.filter(
        auditor=user_ami, pengisian__prodi__fakultas_id=user_ami.fakultas_id,
    ).exists()


@login_required
def user_auditor_list(request):
    if not _is_pengawas(request):
        messages.error(request, 'Halaman ini hanya untuk LP3M/Pimpinan.')
        return render(request, 'ami_user/forbidden.html', {'active_tab': 'user_auditor'})

    role_key = request.GET.get('matrix', 'auditor_de')
    if role_key not in PERMISSION_MATRIX:
        role_key = 'auditor_de'

    base_qs = UserAmi.objects.select_related('user', 'fakultas').order_by('user__username')

    stats = {
        'total_aktif': base_qs.filter(is_aktif=True).count(),
        'auditor_de': base_qs.filter(is_auditor_de=True, is_aktif=True).count(),
        'auditor_visitasi': base_qs.filter(is_auditor_visitasi=True, is_aktif=True).count(),
        'upm': base_qs.filter(is_upm=True, is_aktif=True).count(),
    }
    auditor_qs = base_qs.filter(is_auditor_de=True) | base_qs.filter(is_auditor_visitasi=True)
    auditor_qs = auditor_qs.distinct()
    pakta_signed = auditor_qs.filter(pakta_integritas_signed_at__isnull=False).count()
    pakta_total = auditor_qs.count()

    # --- Filter & pencarian tabel (dimensi terpisah dari `matrix` di atas) ---
    peran_filter = request.GET.get('peran', '')
    status_filter = request.GET.get('status', 'aktif')
    fakultas_filter = request.GET.get('fakultas', '')
    q = request.GET.get('q', '').strip()

    users = base_qs
    if peran_filter == 'auditor_de':
        users = users.filter(is_auditor_de=True)
    elif peran_filter == 'auditor_visitasi':
        users = users.filter(is_auditor_visitasi=True)
    elif peran_filter == 'upm':
        users = users.filter(is_upm=True)
    elif peran_filter == 'lp3m':
        users = users.filter(is_lp3m=True)
    elif peran_filter == 'pimpinan':
        users = users.filter(is_pimpinan=True)
    elif peran_filter == 'auditee':
        users = users.filter(
            is_auditor_de=False, is_auditor_visitasi=False, is_upm=False,
            is_lp3m=False, is_pimpinan=False,
        )

    if status_filter == 'aktif':
        users = users.filter(is_aktif=True)
    elif status_filter == 'nonaktif':
        users = users.filter(is_aktif=False)

    if fakultas_filter:
        users = users.filter(fakultas_id=fakultas_filter)

    if q:
        users = users.filter(
            models.Q(user__first_name__icontains=q) | models.Q(user__last_name__icontains=q) |
            models.Q(user__username__icontains=q) | models.Q(nidn_nip__icontains=q),
        )

    users = list(users)
    for u in users:
        is_auditor = u.is_auditor_de or u.is_auditor_visitasi
        u.konflik = is_auditor and ada_konflik_kepentingan(u)
        u.butuh_pakta = is_auditor and not u.pakta_integritas_signed_at

    paginator = Paginator(users, 12)
    page_obj = paginator.get_page(request.GET.get('page'))

    role_choices = [(key, r['label']) for key, r in PERMISSION_MATRIX.items()]

    return render(request, 'ami_user/user_auditor_list.html', {
        'page_obj': page_obj, 'stats': stats,
        'pakta_signed': pakta_signed, 'pakta_total': pakta_total,
        'role_choices': role_choices, 'role_key': role_key,
        'selected_role': PERMISSION_MATRIX[role_key],
        'fakultas_list': Fakultas.objects.filter(is_aktif=True).order_by('kode'),
        'peran_filter': peran_filter, 'status_filter': status_filter,
        'fakultas_filter': fakultas_filter, 'q': q,
        'active_tab': 'user_auditor',
    })


@login_required
def sign_pakta_integritas(request):
    user_ami = _get_user_ami(request)
    if user_ami is None or not (user_ami.is_auditor_de or user_ami.is_auditor_visitasi):
        messages.error(request, 'Hanya auditor yang perlu menandatangani Pakta Integritas.')
        return redirect('self_assessment:pengisian_detail')

    if request.method == 'POST':
        user_ami.pakta_integritas_signed_at = timezone.now()
        user_ami.save(update_fields=['pakta_integritas_signed_at', 'updated_at'])
        messages.success(request, 'Pakta Integritas berhasil ditandatangani. Terima kasih atas komitmen Anda.')

    if user_ami.is_auditor_de:
        return redirect('de:penugasan_list')
    return redirect('visitasi:visitasi_saya')


@login_required
def user_detail(request, user_id):
    if not _is_pengawas(request):
        messages.error(request, 'Halaman ini hanya untuk LP3M/Pimpinan.')
        return render(request, 'ami_user/forbidden.html', {'active_tab': 'user_auditor'})

    from apps.ami_de.models import DePenugasan

    target = get_object_or_404(UserAmi.objects.select_related('user', 'fakultas', 'prodi'), pk=user_id)
    penugasan_list = DePenugasan.objects.filter(auditor=target).select_related(
        'pengisian__prodi__fakultas', 'siklus',
    ).order_by('-siklus')

    return render(request, 'ami_user/user_detail.html', {
        'target': target,
        'konflik': (target.is_auditor_de or target.is_auditor_visitasi) and ada_konflik_kepentingan(target),
        'penugasan_list': penugasan_list,
        'active_tab': 'user_auditor',
    })


@login_required
def tambah_auditor(request):
    if not _is_pengawas(request):
        messages.error(request, 'Halaman ini hanya untuk LP3M/Pimpinan.')
        return render(request, 'ami_user/forbidden.html', {'active_tab': 'user_auditor'})

    from django.contrib.auth.models import User

    from .forms import TambahAuditorForm

    if request.method == 'POST':
        form = TambahAuditorForm(request.POST)
        if form.is_valid():
            cd = form.cleaned_data
            password = User.objects.make_random_password(length=10)
            user = User.objects.create_user(
                username=cd['username'], email=cd.get('email') or '',
                first_name=cd['first_name'], last_name=cd.get('last_name') or '',
                password=password,
            )
            UserAmi.objects.create(
                user=user, nidn_nip=cd.get('nidn_nip'), fakultas=cd.get('fakultas'), prodi=cd.get('prodi'),
                is_auditor_de=cd.get('is_auditor_de', False), is_auditor_visitasi=cd.get('is_auditor_visitasi', False),
                is_upm=cd.get('is_upm', False), sertifikasi_auditor=cd.get('sertifikasi_auditor'),
                pengalaman_audit_thn=cd.get('pengalaman_audit_thn'),
            )
            messages.success(
                request,
                f'Auditor "{cd["username"]}" berhasil dibuat. Password sementara: {password} '
                '-- catat dan sampaikan ke yang bersangkutan lewat jalur aman (sistem belum '
                'punya integrasi email/SSO untuk mengirim otomatis).',
            )
            return redirect('user_auditor:user_auditor_list')
    else:
        form = TambahAuditorForm()

    return render(request, 'ami_user/tambah_auditor.html', {
        'form': form, 'active_tab': 'user_auditor',
    })


@login_required
def import_excel_auditor(request):
    if not _is_pengawas(request):
        messages.error(request, 'Halaman ini hanya untuk LP3M/Pimpinan.')
        return render(request, 'ami_user/forbidden.html', {'active_tab': 'user_auditor'})

    from django.contrib.auth.models import User

    from .forms import ImportExcelAuditorForm

    hasil = None
    if request.method == 'POST':
        form = ImportExcelAuditorForm(request.POST, request.FILES)
        if form.is_valid():
            import openpyxl

            wb = openpyxl.load_workbook(form.cleaned_data['file'])
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))

            dibuat, diperbarui, gagal = [], [], []
            for i, row in enumerate(rows, start=2):
                if not row or not row[0]:
                    continue
                try:
                    nama, username, email, nidn, fakultas_kode, aud_de, aud_vis, upm, sertifikasi, pengalaman = (
                        (list(row) + [None] * 10)[:10]
                    )
                    if not username:
                        gagal.append(f'Baris {i}: username kosong.')
                        continue
                    fakultas_obj = Fakultas.objects.filter(kode=fakultas_kode).first() if fakultas_kode else None
                    user, user_created = User.objects.get_or_create(
                        username=str(username).strip(),
                        defaults={'first_name': str(nama or username), 'email': str(email or '')},
                    )
                    if user_created:
                        user.set_password(User.objects.make_random_password(length=10))
                        user.save()
                    ua, ua_created = UserAmi.objects.get_or_create(user=user)
                    ua.nidn_nip = str(nidn).strip() if nidn else ua.nidn_nip
                    ua.fakultas = fakultas_obj or ua.fakultas
                    ua.is_auditor_de = str(aud_de).strip().lower() in ('ya', 'yes', 'true', '1') if aud_de is not None else ua.is_auditor_de
                    ua.is_auditor_visitasi = str(aud_vis).strip().lower() in ('ya', 'yes', 'true', '1') if aud_vis is not None else ua.is_auditor_visitasi
                    ua.is_upm = str(upm).strip().lower() in ('ya', 'yes', 'true', '1') if upm is not None else ua.is_upm
                    ua.sertifikasi_auditor = str(sertifikasi).strip() if sertifikasi else ua.sertifikasi_auditor
                    ua.pengalaman_audit_thn = int(pengalaman) if pengalaman else ua.pengalaman_audit_thn
                    ua.save()
                    (dibuat if (user_created or ua_created) else diperbarui).append(f'{username} ({nama or "-"})')
                except Exception as e:
                    gagal.append(f'Baris {i}: {e}')

            hasil = {'dibuat': dibuat, 'diperbarui': diperbarui, 'gagal': gagal}
    else:
        form = ImportExcelAuditorForm()

    return render(request, 'ami_user/import_excel.html', {
        'form': form, 'hasil': hasil, 'active_tab': 'user_auditor',
    })


@login_required
def download_template_excel(request):
    if not _is_pengawas(request):
        messages.error(request, 'Halaman ini hanya untuk LP3M/Pimpinan.')
        return render(request, 'ami_user/forbidden.html', {'active_tab': 'user_auditor'})

    import openpyxl
    from django.http import HttpResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Template Auditor'
    ws.append([
        'Nama', 'Username', 'Email', 'NIDN', 'Kode Fakultas',
        'Auditor DE (Ya/Tidak)', 'Auditor Visitasi (Ya/Tidak)', 'UPM (Ya/Tidak)',
        'Sertifikasi', 'Pengalaman (tahun)',
    ])
    ws.append(['Dr. Contoh Nama', 'contoh.nama', 'contoh@unisan-g.id', '0012345678', 'FEB', 'Ya', 'Tidak', 'Tidak', 'CIIQA', 3])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="template-import-auditor.xlsx"'
    wb.save(response)
    return response


@login_required
def generate_sk(request):
    if not _is_pengawas(request):
        messages.error(request, 'Halaman ini hanya untuk LP3M/Pimpinan.')
        return render(request, 'ami_user/forbidden.html', {'active_tab': 'user_auditor'})

    import io

    from django.http import HttpResponse
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    from apps.ami_core.models import Siklus

    siklus = Siklus.objects.filter(is_current=True).first()
    auditor_de = UserAmi.objects.filter(is_auditor_de=True, is_aktif=True).select_related('user', 'fakultas').order_by('user__first_name')
    auditor_vis = UserAmi.objects.filter(is_auditor_visitasi=True, is_aktif=True).select_related('user', 'fakultas').order_by('user__first_name')
    upm = UserAmi.objects.filter(is_upm=True, is_aktif=True).select_related('user', 'fakultas').order_by('user__first_name')

    document = Document()
    title = document.add_heading('SURAT KEPUTUSAN', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph('Penetapan Tim Auditor Audit Mutu Internal')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if siklus and siklus.sk_rektor_no:
        document.add_paragraph(f'Nomor: {siklus.sk_rektor_no}')
    else:
        document.add_paragraph('Nomor: (belum ditetapkan -- lengkapi lewat Siklus AMI di admin)')
    document.add_paragraph(f'Siklus: {siklus if siklus else "-"}')
    document.add_paragraph(f'Tanggal SK: {siklus.sk_rektor_tgl if siklus and siklus.sk_rektor_tgl else "(belum ditetapkan)"}')
    document.add_paragraph()

    for judul, qs in [('Auditor Desk Evaluasi', auditor_de), ('Auditor Visitasi', auditor_vis), ('UPM Fakultas', upm)]:
        document.add_heading(judul, level=2)
        if qs.exists():
            table = document.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, h in enumerate(['Nama', 'NIDN', 'Fakultas']):
                hdr[i].text = h
                hdr[i].paragraphs[0].runs[0].bold = True
            for a in qs:
                row = table.add_row().cells
                row[0].text = str(a)
                row[1].text = a.nidn_nip or '-'
                row[2].text = str(a.fakultas) if a.fakultas_id else '-'
        else:
            document.add_paragraph('Belum ada yang ditetapkan pada role ini.')
        document.add_paragraph()

    document.add_paragraph()
    ttd = document.add_paragraph()
    ttd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ttd.add_run('Rektor Universitas Ichsan Gorontalo\n\n\n\n( _____________________________ )')

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="sk-tim-auditor-siklus{siklus.no_siklus if siklus else "x"}.docx"'
    return response


@login_required
def konfirmasi_pakta(request, user_id):
    if not _is_pengawas(request):
        messages.error(request, 'Halaman ini hanya untuk LP3M/Pimpinan.')
        return render(request, 'ami_user/forbidden.html', {'active_tab': 'user_auditor'})

    target = get_object_or_404(UserAmi, pk=user_id)

    if request.method == 'POST':
        target.pakta_integritas_signed_at = timezone.now()
        target.save(update_fields=['pakta_integritas_signed_at', 'updated_at'])
        messages.success(
            request,
            f'Pakta Integritas {target} dicatat sebagai sudah diterima (konfirmasi administratif LP3M -- '
            'bukan tanda tangan digital oleh yang bersangkutan).',
        )
        return redirect('user_auditor:user_auditor_list')

    return render(request, 'ami_user/konfirmasi_pakta.html', {
        'target': target, 'active_tab': 'user_auditor',
    })
